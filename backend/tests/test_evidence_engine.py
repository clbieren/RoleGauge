"""
Comprehensive Test Suite for Evidence Engine & Multi-Source Fusion.

Verifies the 4 core mathematical & override scenarios required by RoleGauge:
1. CV-only scenario (capped at self_reported / claimed)
2. GitHub-only scenario (high confidence, evidence_found)
3. Combined (GitHub + CV) scenario (diminishing factor proof: 0.70 + 0.15 * 0.10 = 0.715)
4. Assessment override scenario (verified_gap hard-resets confidence to 0.0)
5. End-to-end POST /api/analyze router integration
"""

import io
import os
import sys
import tempfile
from unittest.mock import AsyncMock, patch

if sys.platform == "win32":
    if isinstance(sys.stdout, io.TextIOWrapper):
        sys.stdout.reconfigure(encoding="utf-8")
    if isinstance(sys.stderr, io.TextIOWrapper):
        sys.stderr.reconfigure(encoding="utf-8")

import pytest
from fastapi.testclient import TestClient

# Add parent directory to path
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app.main import app
from app.services.evidence_detector import EvidenceSignal, SubskillEvidenceResult
from app.services.evidence_engine import EvidenceEngine, UnifiedSubskillEvidence
from app.services.kb_loader import KnowledgeBase, kb


@pytest.fixture(scope="module", autouse=True)
def init_kb():
    """Ensure Knowledge Base is loaded."""
    kb.load()
    return kb


@pytest.fixture
def evidence_engine(init_kb):
    """Create an EvidenceEngine instance for devops role."""
    return EvidenceEngine(init_kb, "devops", "mid")


# ──────────────────────────────────────────────
# 1. Four Core Scenario Proofs
# ──────────────────────────────────────────────

class TestEvidenceEngineScenarios:
    """Test the 4 required mathematical & override proof scenarios."""

    def test_scenario_1_cv_only(self, evidence_engine):
        """
        Scenario 1: CV-only evidence.
        Self-reported CV skill produces claimed status, capped at self_reported ceiling (0.30).
        """
        test_key = "docker.dockerfile"
        cv_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="claimed",
                signals=[
                    EvidenceSignal(
                        source="cv_skills_list",
                        file_path="cv/skills",
                        matched_text="Docker",
                        strength=0.15,
                    )
                ],
            )
        }

        unified = evidence_engine.merge_evidence(cv_skills_evidence=cv_signals)
        result = unified[test_key]

        # Assertions
        assert result.final_confidence == 0.15
        assert result.final_status == "claimed"
        assert "self_reported_skills_and_summary (0.30)" in result.ceiling_applied
        assert "primary=0.15 (cv_skills_list)" in result.calculation_trace
        assert len(result.contributing_sources) == 1
        assert result.contributing_sources[0].source == "cv_skills_list"

        print(f"\n[SCENARIO 1 - CV ONLY] {test_key}:")
        print(f"  Confidence: {result.final_confidence}")
        print(f"  Status: {result.final_status}")
        print(f"  Ceiling: {result.ceiling_applied}")
        print(f"  Trace: {result.calculation_trace}")

    def test_scenario_2_github_only(self, evidence_engine):
        """
        Scenario 2: GitHub-only evidence.
        Verified GitHub code presence produces evidence_found status, github_present ceiling (1.0).
        """
        test_key = "docker.dockerfile"
        github_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="evidence_found",
                signals=[
                    EvidenceSignal(
                        source="file_presence",
                        file_path="Dockerfile",
                        matched_text="Dockerfile present with multi-stage build",
                        strength=0.70,
                    )
                ],
            )
        }

        unified = evidence_engine.merge_evidence(github_evidence=github_signals)
        result = unified[test_key]

        # Assertions
        assert result.final_confidence == 0.70
        assert result.final_status == "evidence_found"
        assert "github_present (1.00)" in result.ceiling_applied
        assert "primary=0.70 (file_presence)" in result.calculation_trace
        assert len(result.contributing_sources) == 1
        assert result.contributing_sources[0].source == "file_presence"

        print(f"\n[SCENARIO 2 - GITHUB ONLY] {test_key}:")
        print(f"  Confidence: {result.final_confidence}")
        print(f"  Status: {result.final_status}")
        print(f"  Ceiling: {result.ceiling_applied}")
        print(f"  Trace: {result.calculation_trace}")

    def test_scenario_3_combined_github_plus_cv_diminishing_proof(self, evidence_engine):
        """
        Scenario 3: Combined (GitHub + CV) evidence with diminishing factor proof.
        Formula: primary (0.70) + secondary (0.15 * 0.10) = 0.715.
        
        Proof points:
        1. final_confidence (0.715) > github_only (0.70)
        2. final_confidence (0.715) < raw_sum (0.70 + 0.15 = 0.85)
        3. diminishing_factor (0.10) mathematically verified
        """
        test_key = "docker.dockerfile"
        github_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="evidence_found",
                signals=[
                    EvidenceSignal(
                        source="file_presence",
                        file_path="Dockerfile",
                        matched_text="Dockerfile present with multi-stage build",
                        strength=0.70,
                    )
                ],
            )
        }
        cv_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="claimed",
                signals=[
                    EvidenceSignal(
                        source="cv_skills_list",
                        file_path="cv/skills",
                        matched_text="Docker listed in skills section",
                        strength=0.15,
                    )
                ],
            )
        }

        unified = evidence_engine.merge_evidence(
            github_evidence=github_signals,
            cv_skills_evidence=cv_signals,
        )
        result = unified[test_key]

        # Exact mathematical assertions
        expected_raw = 0.70 + (0.15 * 0.10)  # 0.715
        assert result.final_confidence == pytest.approx(0.715, abs=1e-4)
        assert result.final_confidence > 0.70, "Combined confidence must exceed GitHub-only (0.70)"
        assert result.final_confidence < 0.85, "Combined confidence must be less than unweighted sum (0.85)"
        assert result.final_status == "evidence_found"
        assert len(result.contributing_sources) == 2

        # Verify trace string format
        assert "primary=0.70" in result.calculation_trace
        assert "secondary=[0.15*0.10]=0.015" in result.calculation_trace
        assert "raw=0.715" in result.calculation_trace
        assert "final=0.715" in result.calculation_trace

        print(f"\n[SCENARIO 3 - COMBINED GITHUB+CV] {test_key}:")
        print(f"  Confidence: {result.final_confidence}")
        print(f"  Status: {result.final_status}")
        print(f"  Diminishing Proof: 0.70 < {result.final_confidence} < 0.85 (exact: 0.70 + 0.15*0.10 = {expected_raw})")
        print(f"  Trace: {result.calculation_trace}")

    def test_scenario_4_assessment_override_verified_gap(self, evidence_engine):
        """
        Scenario 4: Assessment override scenario.
        Injecting a 'verified_gap' assessment failure hard-caps confidence to 0.00
        despite strong GitHub (0.70) and CV (0.15) evidence.
        """
        test_key = "docker.dockerfile"
        github_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="evidence_found",
                signals=[
                    EvidenceSignal(
                        source="file_presence",
                        file_path="Dockerfile",
                        matched_text="Dockerfile present with multi-stage build",
                        strength=0.70,
                    )
                ],
            )
        }
        cv_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="claimed",
                signals=[
                    EvidenceSignal(
                        source="cv_skills_list",
                        file_path="cv/skills",
                        matched_text="Docker listed in skills section",
                        strength=0.15,
                    )
                ],
            )
        }
        assessment_signals = {
            test_key: [
                {
                    "status": "verified_gap",
                    "score": 0.0,
                    "strength": 0.0,
                    "detail": "Candidate failed Docker multi-stage build live coding challenge",
                }
            ]
        }

        unified = evidence_engine.merge_evidence(
            github_evidence=github_signals,
            cv_skills_evidence=cv_signals,
            assessment_evidence=assessment_signals,
        )
        result = unified[test_key]

        # Assessment override assertions
        assert result.final_confidence == 0.0
        assert result.final_status == "verified_gap"
        assert "assessment_override_verified_gap" in result.ceiling_applied
        assert "assessment_override: verified_gap" in result.calculation_trace

        print(f"\n[SCENARIO 4 - ASSESSMENT OVERRIDE] {test_key}:")
        print(f"  Confidence: {result.final_confidence} (hard reset from 0.715)")
        print(f"  Status: {result.final_status}")
        print(f"  Ceiling: {result.ceiling_applied}")
        print(f"  Trace: {result.calculation_trace}")


# ──────────────────────────────────────────────
# 2. Multi-Signal Diminishing & Ceilings
# ──────────────────────────────────────────────

class TestMultiSignalFusion:
    """Test complex multi-signal fusion and ceiling edge cases."""

    def test_three_secondary_signals_diminishing(self, evidence_engine):
        """Primary (0.60) + 3 secondary signals (0.40, 0.30, 0.20) with 0.10 diminishing factor."""
        test_key = "docker.dockerfile"
        github_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="evidence_found",
                signals=[
                    EvidenceSignal(source="file_presence", file_path="Dockerfile", matched_text="Dockerfile", strength=0.60),
                    EvidenceSignal(source="content_match", file_path="compose.yml", matched_text="docker-compose", strength=0.40),
                ],
            )
        }
        cv_signals = {
            test_key: SubskillEvidenceResult(
                composite_key=test_key,
                subskill_name="Dockerfile Creation",
                status="claimed",
                signals=[
                    EvidenceSignal(source="cv_project", file_path="cv/projects", matched_text="Docker containers", strength=0.30),
                    EvidenceSignal(source="cv_skills_list", file_path="cv/skills", matched_text="Docker", strength=0.20),
                ],
            )
        }

        unified = evidence_engine.merge_evidence(
            github_evidence=github_signals,
            cv_skills_evidence=cv_signals,
        )
        result = unified[test_key]

        # Primary = 0.60, Secondary sum = (0.40 + 0.30 + 0.20) * 0.10 = 0.09
        # Raw = 0.69
        expected = 0.60 + (0.90 * 0.10)
        assert result.final_confidence == pytest.approx(expected, abs=1e-4)

    def test_certifications_only_ceiling(self, evidence_engine):
        """Certifications only should be capped at certifications_only ceiling (0.20)."""
        test_key = "docker.dockerfile"
        cert_matches = [
            {
                "certificate_name": "Docker Certified Associate (DCA)",
                "classification": "recognized_relevant",
                "matched_composite_keys": [test_key],
                "strength": 0.50,  # Even with strong cert signal
            }
        ]

        unified = evidence_engine.merge_evidence(cv_cert_matches=cert_matches)
        result = unified[test_key]

        # Capped at 0.20 ceiling
        assert result.final_confidence == 0.20
        assert "certifications_only (0.20)" in result.ceiling_applied
        assert result.final_status == "claimed"


# ──────────────────────────────────────────────
# 3. Router Integration (POST /api/analyze)
# ──────────────────────────────────────────────

class TestAnalyzeRouterIntegration:
    """Test POST /api/analyze router handling GitHub-only, CV-only, and Combined requests."""

    def test_neither_github_nor_cv_returns_400(self):
        """Missing both GitHub username and CV file returns 400."""
        with TestClient(app) as client:
            response = client.post(
                "/api/analyze",
                json={"role_id": "backend", "level": "mid"},
            )
            assert response.status_code == 400
            assert "At least one evidence source" in response.json()["detail"]

    def test_github_only_json_request(self):
        """Pure JSON request with github_username works."""
        with patch("app.routers.analyze.GitHubFetcher.fetch_user_repos", new_callable=AsyncMock) as mock_fetch:
            # Mock repos return
            mock_fetch.return_value = []
            with TestClient(app) as client:
                response = client.post(
                    "/api/analyze",
                    json={
                        "github_username": "testuser",
                        "role_id": "backend",
                        "level": "mid",
                    },
                )
                assert response.status_code == 404  # Expected when 0 repos returned

    def test_cv_only_multipart_request(self):
        """Multipart request with CV file only."""
        from docx import Document
        doc = Document()
        doc.add_paragraph("John Doe\njohn@example.com")
        doc.add_paragraph("Skills\nPython, FastAPI, Docker, PostgreSQL")
        
        docx_io = io.BytesIO()
        doc.save(docx_io)
        docx_io.seek(0)

        with TestClient(app) as client:
            response = client.post(
                "/api/analyze",
                data={"role_id": "backend", "level": "mid"},
                files={"file": ("resume.docx", docx_io, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )
            assert response.status_code == 200
            data = response.json()
            assert data["has_cv"] is True
            assert "readiness_score" in data
            assert len(data["skills"]) > 0

            # Check that calculation_trace is present in subskills
            found_traces = []
            for s in data["skills"]:
                for sub in s["subskills"]:
                    if sub.get("calculation_trace"):
                        found_traces.append(sub["calculation_trace"])
            assert len(found_traces) > 0, "Expected calculation traces in response subskills"


if __name__ == "__main__":
    pytest.main([__file__, "-v", "-s"])
